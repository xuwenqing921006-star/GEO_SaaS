import re
from dataclasses import dataclass
from itertools import chain
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

MAX_FRAGMENT_CHARS = 10_000
MAX_TOTAL_CHARS = 2_000_000
MAX_PDF_PAGES = 500
MAX_SHEET_ROWS = 5_000
MAX_SHEET_COLUMNS = 200
MAX_SHEET_CELLS = 250_000
MAX_WORKSHEETS = 100
PROMPT_INJECTION_PATTERNS = (
    "忽略之前",
    "忽略以上",
    "system prompt",
    "developer message",
    "系统提示词",
    "执行以下指令",
)
PHONE_PATTERN = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
GEO_TEMPLATE_TITLE = "企业资料采集表"
GEO_TEMPLATE_HEADERS = ("序号", "资料项目", "客户填写")


class ParseFailure(RuntimeError):
    pass


class OcrUnavailable(ParseFailure):
    pass


@dataclass(frozen=True)
class ParsedFragment:
    page: int | None
    text: str
    extraction_method: str
    sensitivity_tags: tuple[str, ...]


@dataclass(frozen=True)
class ParseResult:
    fragments: tuple[ParsedFragment, ...]
    page_count: int | None


def _sensitivity_tags(text: str) -> tuple[str, ...]:
    tags: list[str] = []
    lowered = text.casefold()
    if any(pattern in lowered for pattern in PROMPT_INJECTION_PATTERNS):
        tags.append("prompt_injection")
    if PHONE_PATTERN.search(text):
        tags.append("personal_information")
    return tuple(tags)


def _fragments(text: str, *, page: int | None, method: str) -> list[ParsedFragment]:
    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not cleaned:
        return []
    if len(cleaned) > MAX_TOTAL_CHARS:
        raise ParseFailure("extracted text exceeds the safe processing limit")
    return [
        ParsedFragment(
            page=page,
            text=chunk,
            extraction_method=method,
            sensitivity_tags=_sensitivity_tags(chunk),
        )
        for offset in range(0, len(cleaned), MAX_FRAGMENT_CHARS)
        if (chunk := cleaned[offset : offset + MAX_FRAGMENT_CHARS])
    ]


def _enforce_total_limit(fragments: list[ParsedFragment]) -> None:
    if sum(len(fragment.text) for fragment in fragments) > MAX_TOTAL_CHARS:
        raise ParseFailure("extracted text exceeds the safe processing limit")


def _extend_with_limit(
    fragments: list[ParsedFragment],
    text: str,
    *,
    page: int | None,
    method: str,
    total_chars: int,
) -> int:
    extracted = _fragments(text, page=page, method=method)
    next_total = total_chars + sum(len(fragment.text) for fragment in extracted)
    if next_total > MAX_TOTAL_CHARS:
        raise ParseFailure("extracted text exceeds the safe processing limit")
    fragments.extend(extracted)
    return next_total


def _parse_pdf(path: Path) -> ParseResult:
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ParseFailure("encrypted PDF is not supported")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ParseFailure("PDF page count exceeds the safe processing limit")
    fragments: list[ParsedFragment] = []
    total_chars = 0
    for index, page in enumerate(reader.pages, start=1):
        total_chars = _extend_with_limit(
            fragments,
            page.extract_text() or "",
            page=index,
            method="text",
            total_chars=total_chars,
        )
    return ParseResult(tuple(fragments), len(reader.pages))


def _parse_docx(path: Path) -> ParseResult:
    document = Document(str(path))
    blocks: list[str] = []
    total_chars = 0
    for block in chain(
        (paragraph.text for paragraph in document.paragraphs),
        (
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ),
    ):
        total_chars += len(block)
        if total_chars > MAX_TOTAL_CHARS:
            raise ParseFailure("extracted text exceeds the safe processing limit")
        blocks.append(block)
    fragments = _fragments("\n".join(blocks), page=None, method="text")
    _enforce_total_limit(fragments)
    return ParseResult(tuple(fragments), None)


def _parse_xlsx(path: Path) -> ParseResult:
    workbook = load_workbook(path, read_only=True, data_only=True)
    fragments: list[ParsedFragment] = []
    total_chars = 0
    try:
        if len(workbook.worksheets) > MAX_WORKSHEETS:
            raise ParseFailure("spreadsheet contains too many worksheets")
        for sheet in workbook.worksheets:
            if (
                sheet.max_column > MAX_SHEET_COLUMNS
                or sheet.max_row * sheet.max_column > MAX_SHEET_CELLS
            ):
                raise ParseFailure("spreadsheet dimensions exceed the safe processing limit")
            title = str(sheet.cell(row=1, column=1).value or "").strip()
            headers = tuple(
                str(sheet.cell(row=4, column=index).value or "").strip() for index in range(1, 4)
            )
            if title == GEO_TEMPLATE_TITLE and headers == GEO_TEMPLATE_HEADERS:
                for index, row in enumerate(
                    sheet.iter_rows(min_row=5, max_col=3, values_only=True),
                    start=5,
                ):
                    if index > MAX_SHEET_ROWS:
                        raise ParseFailure(
                            "spreadsheet row count exceeds the safe processing limit"
                        )
                    sequence, raw_label, raw_value = row
                    if sequence is None or raw_label is None or raw_value is None:
                        continue
                    if not str(sequence).strip().isdigit():
                        continue
                    label = str(raw_label).strip().lstrip("*").strip()
                    value = str(raw_value).strip()
                    if not label or not value:
                        continue
                    text = f"{label}: {value}"
                    total_chars += len(text)
                    if total_chars > MAX_TOTAL_CHARS:
                        raise ParseFailure("extracted text exceeds the safe processing limit")
                    fragments.append(
                        ParsedFragment(
                            page=None,
                            text=text,
                            extraction_method="geo_template",
                            sensitivity_tags=_sensitivity_tags(text),
                        ),
                    )
                continue
            lines = [f"工作表: {sheet.title}"]
            sheet_chars = len(lines[0])
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if index > MAX_SHEET_ROWS:
                    raise ParseFailure("spreadsheet row count exceeds the safe processing limit")
                values = [str(value).strip() for value in row if value is not None]
                if values:
                    line = " | ".join(values)
                    sheet_chars += len(line) + 1
                    if total_chars + sheet_chars > MAX_TOTAL_CHARS:
                        raise ParseFailure("extracted text exceeds the safe processing limit")
                    lines.append(line)
            total_chars = _extend_with_limit(
                fragments,
                "\n".join(lines),
                page=None,
                method="text",
                total_chars=total_chars,
            )
    finally:
        workbook.close()
    return ParseResult(tuple(fragments), None)


def _parse_text(path: Path) -> ParseResult:
    fragments = _fragments(
        path.read_text(encoding="utf-8-sig"),
        page=None,
        method="text",
    )
    _enforce_total_limit(fragments)
    return ParseResult(tuple(fragments), None)


def _validate_image(path: Path) -> None:
    with Image.open(path) as image:
        image.verify()


def parse_material(path: Path, media_type: str) -> ParseResult:
    if media_type == "application/pdf":
        return _parse_pdf(path)
    if media_type.endswith("wordprocessingml.document"):
        return _parse_docx(path)
    if media_type.endswith("spreadsheetml.sheet"):
        return _parse_xlsx(path)
    if media_type == "text/plain":
        return _parse_text(path)
    if media_type in {"image/png", "image/jpeg"}:
        _validate_image(path)
        raise OcrUnavailable("本机 OCR 尚未配置，请保留原图并手工录入事实")
    raise ParseFailure("unsupported material media type")
