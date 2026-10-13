# ruff: noqa: RUF001

from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document as DocxDocument
from pydantic import SecretStr

from app.config import Settings
from app.content import _export_path, _render_export
from app.database import ContentExport, ContentVersion, FactSnapshot


def build_export_fixture() -> tuple[ContentVersion, FactSnapshot]:
    document_id = uuid4()
    snapshot = FactSnapshot(
        id=uuid4(),
        client_id=uuid4(),
        brief_id=uuid4(),
        brief_version=2,
        facts=[
            {
                "id": "fact-1",
                "key": "standard_name",
                "label": "企业标准名称",
                "value": "莱恩美术",
                "category": "identity",
            },
        ],
        evidence=[
            {
                "id": "evidence-1",
                "fact_id": "fact-1",
                "source_type": "web",
                "source_label": "公开机构介绍",
                "source_url": "https://example.com/lion",
                "quote": "莱恩美术",
                "observed_at": "2026-07-27",
            },
        ],
        brief_context={
            "title": "即墨少儿美术怎么选",
            "content_type": "selection_guide",
            "target_keywords": ["即墨少儿美术"],
            "retest_questions": ["即墨少儿美术培训怎么选？"],
            "target_level": "L2",
            "audience": "本地家长",
            "angle": "选择条件",
            "suggested_channel_ids": ["zhihu", "baijiahao"],
            "expected_recommendation_reason": "实体与服务关系清楚",
        },
        prohibited_claims=[],
    )
    version = ContentVersion(
        id=uuid4(),
        document_id=document_id,
        version_number=3,
        title="即墨少儿美术怎么选",
        body_markdown="# 旧标题\n\n莱恩美术的公开信息说明。",
        outline_markdown=None,
        change_type="generated",
        provider="rules",
        model="fixed-mock-v1",
        fact_snapshot_id=snapshot.id,
        prompt_version="content-v1",
    )
    return version, snapshot


def test_markdown_export_has_title_facts_sources_without_internal_geo_guidance() -> None:
    version, snapshot = build_export_fixture()

    rendered = _render_export(
        version,
        snapshot,
        format_name="markdown",
        include_sources=True,
    ).decode("utf-8")

    assert rendered.startswith("# 即墨少儿美术怎么选\n")
    assert "# 旧标题" not in rendered
    assert "## 已确认企业事实" in rendered
    assert "企业标准名称：莱恩美术" in rendered
    assert "## 事实来源" in rendered
    assert "https://example.com/lion" in rendered
    assert "摘录：莱恩美术" in rendered
    assert "发布与验证建议" not in rendered
    assert "目标关键词" not in rendered
    assert "AI 平台" not in rendered


def test_docx_export_preserves_same_business_sections() -> None:
    version, snapshot = build_export_fixture()

    payload = _render_export(
        version,
        snapshot,
        format_name="docx",
        include_sources=True,
    )
    document = DocxDocument(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert document.paragraphs[0].text == "即墨少儿美术怎么选"
    assert "已确认企业事实" in text
    assert "公开机构介绍" in text
    assert "发布与验证建议" not in text
    assert "AI 平台" not in text


def test_export_path_is_scoped_below_storage_root(tmp_path: Path) -> None:
    version, _ = build_export_fixture()
    export = ContentExport(
        id=uuid4(),
        document_id=version.document_id,
        content_version_id=version.id,
        format="markdown",
        include_sources=True,
    )
    settings = Settings(
        app_env="test",
        db_password=SecretStr("unused"),
        storage_root=tmp_path,
        model_execution_mode="fixed_mock",
    )

    path = _export_path(settings, export)

    assert (tmp_path.resolve() / "exports") in path.parents
    assert path.name == f"{export.id}.md"
