# ruff: noqa: RUF001

from collections.abc import Iterator
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import cast
from uuid import UUID, uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from pydantic import SecretStr
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.config import Settings
from app.database import Evidence, Fact, FactConflict, Job, Material
from app.jobs import process_one_job
from tests.auth.test_auth_api import build_client
from tests.clients.test_clients_api import client_payload


@pytest.fixture
def prepared_client(tmp_path: Path) -> Iterator[TestClient]:
    client = build_client(tmp_path)
    yield client
    client.close()


@pytest.mark.integration
@pytest.mark.asyncio
async def test_t007_api_journey_runs_in_one_async_context(
    tmp_path: Path,
    prepared_client: TestClient,
) -> None:
    sync_client = prepared_client
    application = cast(FastAPI, sync_client.app)
    factory = cast(
        async_sessionmaker[AsyncSession],
        application.state.session_factory,
    )
    transport = ASGITransport(app=application)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        login = await client.post(
            "/api/v1/auth/login",
            json={"username": "owner", "password": "CorrectPassword2026!"},
        )
        assert login.status_code == 200
        csrf = login.json()["data"]["csrf_token"]
        headers = {"X-CSRF-Token": csrf}

        created = await client.post(
            "/api/v1/clients",
            headers=headers,
            json=client_payload(),
        )
        assert created.status_code == 201
        client_id = created.json()["data"]["id"]

        uploaded = await client.post(
            f"/api/v1/clients/{client_id}/materials",
            headers=headers,
            data={"category": "service", "authorization_status": "authorized"},
            files={
                "files": (
                    "service.txt",
                    "核心服务: 家电维修\n服务范围: 即墨区".encode(),
                    "text/plain",
                ),
            },
        )
        assert uploaded.status_code == 201
        material = uploaded.json()["data"]["items"][0]
        listed = await client.get(
            f"/api/v1/clients/{client_id}/materials",
            params={"category": "service", "extraction_status": "not_started"},
        )
        assert listed.json()["data"]["total"] == 1

        job_key = str(uuid4())
        accepted = await client.post(
            f"/api/v1/materials/{material['id']}/extract",
            headers={**headers, "Idempotency-Key": job_key},
            json={"force": False, "languages": ["zh-CN"]},
        )
        assert accepted.status_code == 202
        repeated = await client.post(
            f"/api/v1/materials/{material['id']}/extract",
            headers={**headers, "Idempotency-Key": job_key},
            json={"force": False, "languages": ["zh-CN"]},
        )
        assert repeated.json()["data"]["job_id"] == accepted.json()["data"]["job_id"]
        assert (
            await client.get(f"/api/v1/jobs/{accepted.json()['data']['job_id']}")
        ).status_code == 200

        settings = Settings(
            app_env="test",
            storage_root=tmp_path,
            db_password=SecretStr("unused"),
        )
        assert await process_one_job(factory, settings) is True
        detail = await client.get(f"/api/v1/materials/{material['id']}")
        assert detail.json()["data"]["material"]["extraction_status"] == "succeeded"

        facts = await client.get(
            f"/api/v1/clients/{client_id}/facts",
            params={"category": "service", "status": "client_provided", "q": "家电"},
        )
        candidate = facts.json()["data"]["items"][0]
        evidence_id = candidate["evidence"][0]["id"]
        confirmed = await client.post(
            f"/api/v1/facts/{candidate['id']}/confirm",
            headers=headers,
            json={
                "version": candidate["version"],
                "evidence_ids": [evidence_id],
                "confirmation_note": "核验原资料并确认公开",
            },
        )
        assert confirmed.status_code == 200
        reopened = await client.patch(
            f"/api/v1/facts/{candidate['id']}",
            headers=headers,
            json={
                "version": confirmed.json()["data"]["version"],
                "value": "提供家电与水电维修",
            },
        )
        assert reopened.json()["data"]["status"] == "client_provided"
        linked = await client.post(
            f"/api/v1/facts/{candidate['id']}/evidence",
            headers=headers,
            json={
                "source_type": "material",
                "material_id": material["id"],
                "source_label": "服务资料补充段落",
                "source_url": None,
                "quote": "提供家电与水电维修",
                "observed_at": "2026-07-27",
                "authorization_status": "authorized",
            },
        )
        assert linked.status_code == 201

        profile = await client.get(f"/api/v1/clients/{client_id}/knowledge-profile")
        profile_id = str(uuid4())
        saved_profile = await client.patch(
            f"/api/v1/clients/{client_id}/knowledge-profile",
            headers=headers,
            json={
                "version": profile.json()["data"]["version"],
                "services": [
                    {
                        "id": profile_id,
                        "name": "家电维修",
                        "description": "家庭常见家电检测维修",
                        "target_audience": "即墨区家庭用户",
                        "price_statement": "按检测结果报价",
                        "limitations": ["不提供高空作业"],
                        "active": True,
                    },
                ],
                "capabilities": [],
                "prohibited_claims": [
                    {
                        "id": str(uuid4()),
                        "phrase": "保证一次修好",
                        "reason": "无法验证的结果承诺",
                        "source": "user",
                    },
                ],
                "external_presences": [],
            },
        )
        assert saved_profile.status_code == 200
        updated_profile = await client.patch(
            f"/api/v1/clients/{client_id}/knowledge-profile",
            headers=headers,
            json={
                "version": saved_profile.json()["data"]["version"],
                "services": saved_profile.json()["data"]["services"],
                "capabilities": [],
                "prohibited_claims": saved_profile.json()["data"]["prohibited_claims"],
                "external_presences": [],
            },
        )
        assert updated_profile.status_code == 200

        async with factory() as database:
            conflict_fact = Fact(
                client_id=UUID(client_id),
                category="location",
                key="location.conflict_test",
                label="服务范围",
                value="即墨区",
                status="conflict",
                confidence=50,
            )
            conflict_evidence = Evidence(
                client_id=UUID(client_id),
                material_id=UUID(material["id"]),
                source_type="material",
                source_label="服务资料",
                quote="服务范围: 即墨区",
                observed_at=date(2026, 7, 27),
                authorization_status="authorized",
            )
            database.add_all([conflict_fact, conflict_evidence])
            await database.flush()
            conflict = FactConflict(
                client_id=UUID(client_id),
                fact_key=conflict_fact.key,
                label=conflict_fact.label,
                candidates=[
                    {
                        "value": "即墨区",
                        "source_label": "服务资料",
                        "evidence_id": str(conflict_evidence.id),
                    },
                ],
            )
            database.add(conflict)
            await database.commit()
            conflict_id = str(conflict.id)
            conflict_version = conflict.version

        conflict_list = await client.get(
            f"/api/v1/clients/{client_id}/conflicts",
            params={"status": "open"},
        )
        assert conflict_list.json()["data"]["items"]
        resolved = await client.post(
            f"/api/v1/conflicts/{conflict_id}/resolve",
            headers=headers,
            json={
                "version": conflict_version,
                "action": "select",
                "selected_evidence_id": str(conflict_evidence.id),
                "manual_value": None,
                "resolution_note": "采用公开资料",
            },
        )
        assert resolved.status_code == 200

        async with factory() as database:
            failed = Job(
                client_id=UUID(client_id),
                material_id=UUID(material["id"]),
                job_type="material_extract",
                status="failed",
                attempts=1,
                idempotency_key=str(uuid4()),
                payload={"force": False, "languages": ["zh-CN"]},
                error_message="临时失败",
            )
            database.add(failed)
            stored_material = await database.get(Material, UUID(material["id"]))
            assert stored_material is not None
            stored_material.extraction_status = "failed"
            await database.commit()
            failed_id = str(failed.id)

        retry_key = str(uuid4())
        retried = await client.post(
            f"/api/v1/jobs/{failed_id}/retry",
            headers={**headers, "Idempotency-Key": retry_key},
            json={"reason": "修复后重试"},
        )
        assert retried.status_code == 202
        idempotent_retry = await client.post(
            f"/api/v1/jobs/{failed_id}/retry",
            headers={**headers, "Idempotency-Key": retry_key},
            json={"reason": "修复后重试"},
        )
        assert idempotent_retry.json()["data"]["job_id"] == retried.json()["data"]["job_id"]
        cancelled = await client.post(
            f"/api/v1/jobs/{retried.json()['data']['job_id']}/cancel",
            headers=headers,
            json={"reason": "本次不再解析"},
        )
        assert cancelled.status_code == 200


@pytest.mark.integration
@pytest.mark.asyncio
async def test_simple_customer_template_import_enables_diagnostic(
    tmp_path: Path,
    prepared_client: TestClient,
) -> None:
    application = cast(FastAPI, prepared_client.app)
    factory = cast(
        async_sessionmaker[AsyncSession],
        application.state.session_factory,
    )
    transport = ASGITransport(app=application)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        login_response = await client.post(
            "/api/v1/auth/login",
            json={"username": "owner", "password": "CorrectPassword2026!"},
        )
        csrf = login_response.json()["data"]["csrf_token"]
        headers = {"X-CSRF-Token": csrf}
        created = await client.post(
            "/api/v1/clients",
            headers=headers,
            json=client_payload(),
        )
        client_id = created.json()["data"]["id"]

        stream = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet["A1"] = "企业资料采集表"
        for column, value in enumerate(
            ("序号", "资料项目", "客户填写", "填写说明 / 示例"),
            start=1,
        ):
            sheet.cell(row=4, column=column, value=value)
        rows = [
            (1, "* 企业或机构标准名称", "莱恩美术"),
            (4, "* 所在省、市、区县", "山东省青岛市即墨区"),
            (5, "* 详细地址 / 门店地址", "德馨校区：即墨区示例路 1 号"),
            (8, "* 核心服务或产品", "美术：少儿美术培训"),
            (10, "* 主要服务对象", "适合：4—12岁儿童"),
            (11, "* 服务区域 / 上门范围", "即墨城区及周边街道"),
        ]
        for row_number, row in enumerate(rows, start=5):
            for column, value in enumerate(row, start=1):
                sheet.cell(row=row_number, column=column, value=value)
        workbook.save(stream)
        workbook.close()

        uploaded = await client.post(
            f"/api/v1/clients/{client_id}/materials",
            headers=headers,
            data={"category": "other", "authorization_status": "authorized"},
            files={
                "files": (
                    "GEO客户资料采集表模板.xlsx",
                    stream.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )
        assert uploaded.status_code == 201
        material_id = uploaded.json()["data"]["items"][0]["id"]
        accepted = await client.post(
            f"/api/v1/materials/{material_id}/extract",
            headers={**headers, "Idempotency-Key": str(uuid4())},
            json={"force": False, "languages": ["zh-CN"]},
        )
        assert accepted.status_code == 202
        settings = Settings(
            app_env="test",
            storage_root=tmp_path,
            db_password=SecretStr("unused"),
        )
        assert await process_one_job(factory, settings) is True

        facts = await client.get(
            f"/api/v1/clients/{client_id}/facts",
            params={"page_size": 100},
        )
        fact_map = {item["key"]: item["value"] for item in facts.json()["data"]["items"]}
        assert fact_map["identity.standard_name"] == "莱恩美术"
        assert fact_map["location.administrative_area"] == "山东省青岛市即墨区"
        assert fact_map["location.address"] == "德馨校区：即墨区示例路 1 号"
        assert fact_map["location.service_area"] == "即墨城区及周边街道"
        assert fact_map["service.primary"] == "美术：少儿美术培训"
        assert fact_map["audience.primary"] == "适合：4—12岁儿童"

        async with factory() as database:
            service_fact = await database.scalar(
                select(Fact).where(
                    Fact.client_id == UUID(client_id),
                    Fact.key == "service.primary",
                ),
            )
            assert service_fact is not None
            previous_version = service_fact.version
            service_fact.status = "confirmed"
            service_fact.confirmation_note = "历史确认记录"
            await database.commit()

        update_stream = BytesIO()
        update_workbook = Workbook()
        update_sheet = update_workbook.active
        assert update_sheet is not None
        update_sheet["A1"] = "企业资料采集表"
        for column, value in enumerate(
            ("序号", "资料项目", "客户填写", "填写说明 / 示例"),
            start=1,
        ):
            update_sheet.cell(row=4, column=column, value=value)
        update_sheet.append([8, "* 核心服务或产品", "少儿美术课程与创意绘画", ""])
        update_workbook.save(update_stream)
        update_workbook.close()
        updated_upload = await client.post(
            f"/api/v1/clients/{client_id}/materials",
            headers=headers,
            data={"category": "other", "authorization_status": "authorized"},
            files={
                "files": (
                    "GEO客户资料更新.xlsx",
                    update_stream.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )
        updated_material_id = updated_upload.json()["data"]["items"][0]["id"]
        update_accepted = await client.post(
            f"/api/v1/materials/{updated_material_id}/extract",
            headers={**headers, "Idempotency-Key": str(uuid4())},
            json={"force": False, "languages": ["zh-CN"]},
        )
        assert update_accepted.status_code == 202
        assert await process_one_job(factory, settings) is True

        async with factory() as database:
            service_facts = (
                await database.scalars(
                    select(Fact).where(
                        Fact.client_id == UUID(client_id),
                        Fact.key == "service.primary",
                    ),
                )
            ).all()
            assert len(service_facts) == 1
            assert service_facts[0].value == "少儿美术课程与创意绘画"
            assert service_facts[0].status == "client_provided"
            assert service_facts[0].confirmation_note is None
            assert service_facts[0].version == previous_version + 1

        diagnostic = await client.post(
            f"/api/v1/clients/{client_id}/diagnostics",
            headers={**headers, "Idempotency-Key": str(uuid4())},
            json={
                "scope": ["entity", "evidence", "content", "source", "visibility"],
                "monitoring_window_days": 30,
            },
        )
        assert diagnostic.status_code == 202
        assert await process_one_job(factory, settings) is True
        latest = await client.get(
            f"/api/v1/clients/{client_id}/diagnostics/latest",
        )
        assert latest.json()["data"]["status"] == "succeeded"
